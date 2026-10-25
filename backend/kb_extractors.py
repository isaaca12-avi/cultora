"""Extracción de texto desde archivos subidos por admins para la biblioteca del motor."""
from __future__ import annotations
import io
import logging
import re
from typing import Optional, Tuple

SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
MAX_SINGLE_FILE_BYTES = 10 * 1024 * 1024   # 10 MB
MAX_TOTAL_BYTES = 30 * 1024 * 1024         # 30 MB

logger = logging.getLogger(__name__)

try:
    from langdetect import detect_langs, DetectorFactory
    DetectorFactory.seed = 42  # resultado determinista
    _LANGDETECT = True
except ImportError:
    _LANGDETECT = False
    logger.warning("langdetect no disponible — la detección de idioma se salta")


def _clean(text: str) -> str:
    if not text:
        return ""
    # Normalizar saltos de línea y espacios repetidos
    t = text.replace("\r\n", "\n").replace("\r", "\n")
    t = re.sub(r"[ \t]+\n", "\n", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise RuntimeError("pypdf no disponible") from e
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return _clean("\n\n".join(parts))


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError("python-docx no disponible") from e
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # Extraer tablas también
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return _clean("\n".join(parts))


def _extract_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return _clean(data.decode(enc))
        except UnicodeDecodeError:
            continue
    return _clean(data.decode("utf-8", errors="ignore"))


def detect_extension(filename: str) -> str:
    if not filename:
        return ""
    name = filename.lower()
    idx = name.rfind(".")
    return name[idx:] if idx >= 0 else ""


def extract_from_bytes(filename: str, data: bytes) -> Tuple[str, str]:
    """Devuelve (text, extension). Lanza ValueError si el formato no está soportado."""
    ext = detect_extension(filename)
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"Formato no soportado: {ext or 'sin extensión'}")
    if len(data) == 0:
        return "", ext
    if ext == ".pdf":
        return _extract_pdf(data), ext
    if ext == ".docx":
        return _extract_docx(data), ext
    # .txt / .md
    return _extract_text(data), ext


def humanize_size(nbytes: int) -> str:
    if nbytes < 1024:
        return f"{nbytes} B"
    if nbytes < 1024 * 1024:
        return f"{nbytes / 1024:.1f} KB"
    return f"{nbytes / (1024 * 1024):.1f} MB"


# Idiomas comunes que probablemente aparezcan en un corpus de gestión cultural
_LANG_LABELS = {
    "es": "Español", "en": "Inglés", "pt": "Portugués", "fr": "Francés",
    "it": "Italiano", "ca": "Catalán", "gl": "Gallego", "eu": "Euskera",
    "de": "Alemán", "nl": "Neerlandés",
}
_LANG_FLAGS = {
    "es": "🇪🇸", "en": "🇬🇧", "pt": "🇵🇹", "fr": "🇫🇷", "it": "🇮🇹",
    "ca": "🇪🇸", "gl": "🇪🇸", "eu": "🇪🇸", "de": "🇩🇪", "nl": "🇳🇱",
}


def detect_language(text: str, min_chars: int = 40) -> Optional[dict]:
    """Detecta el idioma dominante. Devuelve {code, label, flag, confidence} o None."""
    if not _LANGDETECT:
        return None
    sample = (text or "").strip()
    if len(sample) < min_chars:
        return None
    # Tomar una muestra representativa (primeros 4000 chars) para eficiencia
    sample = sample[:4000]
    try:
        candidates = detect_langs(sample)
    except Exception:
        return None
    if not candidates:
        return None
    top = candidates[0]
    code = str(top.lang).lower()
    return {
        "code": code,
        "label": _LANG_LABELS.get(code, code.upper()),
        "flag": _LANG_FLAGS.get(code, "🌐"),
        "confidence": round(float(top.prob), 3),
    }
